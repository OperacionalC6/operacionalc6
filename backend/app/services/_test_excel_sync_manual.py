"""
Teste manual (não faz parte da suíte formal) do módulo excel_sync.py contra
um workbook sintético que reproduz a estrutura REAL descoberta na inspeção do
Construcao.xlsx do usuário (2026-09-24) — sem depender do RPA de verdade
(baixar_looker_bruto é substituído por uma função fake).

Roda com: python -m app.services._test_excel_sync_manual
"""

import sys
from copy import copy
from datetime import date, datetime

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, ".")

import app.services.excel_sync as excel_sync
from app.services.excel_sync import (
    AtualizacaoRecusada,
    arrastar_base_final,
    atualizar_db_apuracaoavista,
    atualizar_db_mercado,
    atualizar_db_pagasanalitico,
)


def montar_workbook_sintetico() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # --- db_pagasanalitico: 12 col fórmula (A:L) + col brutas a partir de M ---
    ws = wb.create_sheet("db_pagasanalitico")
    formula_cols = ["ANO", "MES", "CHAVE_CONTRATO", "CHAVE_AREA_EHS", "AREA_LOJA_EHS", "FILIAL_CONTRATO",
                     "CHAVE_LOJA", "CNPJ_LOJA", "CODIGO_LOJA", "NOME_LOJA", "PRODUTO EGV", "SEGURO"]
    raw_cols = ["ID Proposta", "Dt Relatório", "Lojista", "Status Proposta", "Cd Contrato", "Vl Financiamento"]
    header = formula_cols + raw_cols
    ws.append(header)
    ws.append([
        "=YEAR(N2)", "=MONTH(N2)", '=AC2&"."&S2', "=1", "=2", "=3", "=4", "=5", "=6", "=7", "=8", "=9",
        1001, datetime(2026, 9, 20), "10 - X - 111", "PROPOSTA PAGA", "AU001", 1000.0,
    ])
    # Formatação real (moeda + fonte em negrito + fundo colorido) na
    # linha-modelo — pra testar se a linha nova herda o ESTILO completo, não
    # só o valor (achado real em 2026-09-24: 1ª correção só copiava
    # number_format, usuário reportou que fonte/cor ainda ficavam default).
    col_vl_financiamento = header.index("Vl Financiamento") + 1
    cel_modelo = ws.cell(row=2, column=col_vl_financiamento)
    cel_modelo.number_format = "R$ #,##0.00"
    cel_modelo.font = Font(bold=True, color="FF0000")
    cel_modelo.fill = PatternFill("solid", fgColor="FFFF00")
    # Mesma ideia numa coluna de FÓRMULA (ANO, col 1) — testa se `_arrastar_linha`
    # também copia estilo, não só `_escrever_linhas_brutas`.
    ws.cell(row=2, column=1).fill = PatternFill("solid", fgColor="00FF00")

    # --- db_apuracaoavista: só colunas brutas ---
    ws = wb.create_sheet("db_apuracaoavista")
    ws.append(["Anomes Apuracao", "Cd Contrato", "Status Contrato", "Lojista", "R$ Principal Total"])
    ws.append(["202608", "AU000", "Ativo", "10 - X - 111", 5000.0])

    # --- db_mercado: 6 col fórmula (A:F) + brutas a partir de G ---
    ws = wb.create_sheet("db_mercado")
    formula_cols_m = ["DATA1", "DATA2", "CHAVE_LOJA", "AREA_LOJA_EHS", "COD_LOJA", "GRUPO LOJA"]
    raw_cols_m = ["Nome Gp", "CNPJ Loja", "Mês", "Produção C6"]
    ws.append(formula_cols_m + raw_cols_m)
    ws.append(["=YEAR(I2)", "=MONTH(I2)", "=1", "=2", "=3", "=4", "BRUNO", "111", datetime(2026, 8, 1), 100.0])

    # --- base_final: header em 2 linhas, dados a partir da linha 3 ---
    ws = wb.create_sheet("base_final")
    ws.append(["DATA", "CONTRATO", "VALOR_FINANC"])
    ws.append(["ANO", "COD_CONTRATO", "VALOR_FINANCIAMENTO_R$"])
    ws.append(["=YEAR(B3)", "=db_apuracaoavista!B2", "=db_apuracaoavista!E2"])

    for nome in ("db_carterizacao", "config_carteira", "config_AjustesContrato", "config_GNs"):
        wb.create_sheet(nome)

    return wb


def fake_baixar(report_name, tile_key, *, filter_query_override=None, filter_value_override=None, sessao=None):
    if report_name == "acompanhamento_veiculos":
        # "Vl Financiamento" vem como TEXTO no formato do Looker (ex.: "R$
        # 2,000.00"), igual ao CSV real baixado em 2026-09-24 — reproduz o
        # bug real (linha nova ficava com string em vez de número, diferente
        # das linhas antigas, que sempre tiveram float de verdade).
        return pd.DataFrame({
            "ID Proposta": [2001, 2002],
            "Dt Relatório": [datetime(2026, 9, 24), datetime(2026, 9, 24)],
            "Lojista": ["20 - Y - 222", "20 - Y - 222"],
            "Status Proposta": ["PROPOSTA PAGA", "PROPOSTA APROVADA"],
            "Cd Contrato": ["AU100", "AU101"],
            "Vl Financiamento": ["R$ 2,000.00", "R$ 3,000.00"],
        })
    if report_name == "comissao_avista":
        # Reproduz o bug real encontrado em 2026-09-24: o CSV baixado do Looker
        # tinha uma linha de rodapé/total com "Anomes Apuracao" vazio, o que
        # promove a coluna inteira pra float64 (202609 -> 202609.0) — daí um
        # `.astype(str)` ingênuo gerava "202609.0" e nunca batia com "202609"
        # (430 linhas baixadas, 0 batendo). Aqui simulamos exatamente isso:
        # valores float com .0 de sobra, mais uma linha de rodapé com NaN.
        return pd.DataFrame({
            "Anomes Apuracao": [202609.0, 202609.0, float("nan")],
            "Cd Contrato": ["AU200", "AU201", None],
            "Status Contrato": ["Ativo", "Ativo", None],
            "Lojista": ["30 - Z - 333", "30 - Z - 333", None],
            "R$ Principal Total": [7000.0, 8000.0, 15000.0],
        })
    if report_name == "painel_visita_mercado":
        return pd.DataFrame({
            "Nome Gp": ["BRUNO"],
            "CNPJ Loja": ["333"],
            "Mês": [datetime(2026, 9, 1)],
            "Produção C6": [500.0],
        })
    raise AssertionError(f"relatório inesperado no teste: {report_name}")


def main():
    excel_sync.baixar_looker_bruto = fake_baixar

    wb = montar_workbook_sintetico()

    print("== atualizar_db_pagasanalitico (dia 2026-09-24) ==")
    res = atualizar_db_pagasanalitico(wb, date(2026, 9, 24))
    print(res)
    ws = wb["db_pagasanalitico"]
    assert ws.max_row == 4, f"esperava 4 linhas (header+3 dados), veio {ws.max_row}"
    assert ws.cell(row=3, column=13).value == 2001, "ID Proposta da linha nova errado"
    formula_ano_linha3 = ws.cell(row=3, column=1).value
    assert formula_ano_linha3 == "=YEAR(N3)", f"fórmula ANO não arrastou certo: {formula_ano_linha3!r}"
    formula_chave_linha4 = ws.cell(row=4, column=3).value
    assert formula_chave_linha4 == '=AC4&"."&S4', f"fórmula CHAVE_CONTRATO não arrastou certo: {formula_chave_linha4!r}"
    # `cell.fill` devolve um StyleProxy sem __eq__ de verdade (compara sempre
    # False mesmo com conteúdo idêntico) — copy() desembrulha pro PatternFill
    # real, que aí sim compara por valor.
    assert copy(ws.cell(row=3, column=1).fill) == copy(ws.cell(row=2, column=1).fill), (
        "cor de fundo da coluna de fórmula (ANO) não foi copiada pra linha nova"
    )
    print("  OK: linhas inseridas, fórmulas arrastadas e estilo da coluna de fórmula copiado.\n")

    print("== conferindo normalização de valor monetário + estilo completo (bug real 2026-09-24) ==")
    col_vl_financiamento = 18  # "Vl Financiamento" é a 6ª coluna bruta, após 12 de fórmula (12+6=18)
    cel_modelo = ws.cell(row=2, column=col_vl_financiamento)
    cel_linha3 = ws.cell(row=3, column=col_vl_financiamento)
    valor_linha3 = cel_linha3.value
    assert isinstance(valor_linha3, float), (
        f"'Vl Financiamento' devia ter virado float (era 'R$ 2,000.00' baixado como texto), "
        f"veio {type(valor_linha3).__name__}: {valor_linha3!r}"
    )
    assert valor_linha3 == 2000.0, f"valor convertido errado: {valor_linha3!r}"
    assert cel_linha3.number_format == cel_modelo.number_format, (
        f"number_format não foi copiado: modelo={cel_modelo.number_format!r}, linha nova={cel_linha3.number_format!r}"
    )
    assert copy(cel_linha3.font) == copy(cel_modelo.font), (
        f"fonte não foi copiada: modelo={cel_modelo.font!r}, linha nova={cel_linha3.font!r}"
    )
    assert copy(cel_linha3.fill) == copy(cel_modelo.fill), (
        f"cor de fundo não foi copiada: modelo={cel_modelo.fill!r}, linha nova={cel_linha3.fill!r}"
    )
    print(
        f"  OK: 'R$ 2,000.00' virou {valor_linha3!r} (float), e number_format/fonte/fundo "
        "copiados da linha-modelo.\n"
    )

    print("== atualizar_db_apuracaoavista (mês 202609) ==")
    res = atualizar_db_apuracaoavista(wb, "202609")
    print(res)
    ws = wb["db_apuracaoavista"]
    assert ws.max_row == 4, f"esperava 4 linhas (header + 1 antiga + 2 novas), veio {ws.max_row}"
    assert ws.cell(row=2, column=2).value == "AU000", "linha antiga (202608) não deveria ter sido tocada"
    assert ws.cell(row=3, column=2).value == "AU200", "1ª linha nova errada"
    assert ws.cell(row=4, column=2).value == "AU201", "2ª linha nova errada"
    print("  OK: linha antiga preservada, linhas novas no lugar certo.\n")

    print("== arrastar_base_final ==")
    res = arrastar_base_final(wb)
    print(res)
    ws_bf = wb["base_final"]
    assert ws_bf.max_row == 5, f"esperava base_final ir até a linha 5 (2 header + 3 dados), veio {ws_bf.max_row}"
    f_contrato_l4 = ws_bf.cell(row=4, column=2).value
    f_contrato_l5 = ws_bf.cell(row=5, column=2).value
    assert f_contrato_l4 == "=db_apuracaoavista!B3", f"referência posicional errada (linha 4): {f_contrato_l4!r}"
    assert f_contrato_l5 == "=db_apuracaoavista!B4", f"referência posicional errada (linha 5): {f_contrato_l5!r}"
    print("  OK: base_final acompanhou db_apuracaoavista com o deslocamento certo (2 linhas novas).\n")

    print("== atualizar_db_mercado (mês 202609) ==")
    res = atualizar_db_mercado(wb, "202609")
    print(res)
    ws = wb["db_mercado"]
    assert ws.max_row == 3
    assert res["check_quantidade_bate"] is True
    print("  OK.\n")

    print("== testando recusa (mês fora de ordem em db_apuracaoavista) ==")
    try:
        atualizar_db_apuracaoavista(wb, "202601")
        raise AssertionError("deveria ter recusado!")
    except AtualizacaoRecusada as exc:
        print("  OK, recusou como esperado:", exc)

    print("\n== testando SUBSTITUIÇÃO (rodar o mesmo dia de novo, com dado diferente) ==")

    def fake_baixar_v2(report_name, tile_key, **kwargs):
        return pd.DataFrame({
            "ID Proposta": [3001, 3002, 3003],
            "Dt Relatório": [datetime(2026, 9, 24)] * 3,
            "Lojista": ["40 - W - 444"] * 3,
            "Status Proposta": ["PROPOSTA PAGA", "PROPOSTA PAGA", "PROPOSTA REPROVADA"],
            "Cd Contrato": ["AU300", "AU301", "AU302"],
            "Vl Financiamento": [500.0, 600.0, 700.0],
        })

    excel_sync.baixar_looker_bruto = fake_baixar_v2
    linhas_antes = wb["db_pagasanalitico"].max_row
    res = atualizar_db_pagasanalitico(wb, date(2026, 9, 24))
    print(res)
    ws = wb["db_pagasanalitico"]
    linhas_depois = ws.max_row
    assert res["linhas_baixadas"] == 3
    assert linhas_depois == linhas_antes + 1, (
        f"esperava +1 linha no total (2 antigas removidas, 3 novas inseridas: -2+3=+1), "
        f"veio antes={linhas_antes} depois={linhas_depois}"
    )
    ultimo_id = ws.cell(row=linhas_depois, column=13).value
    assert ultimo_id == 3003, f"última linha deveria ser a proposta 3003 (substituição), veio {ultimo_id}"
    penultimo_id = ws.cell(row=linhas_depois - 2, column=13).value
    assert penultimo_id == 3001, f"linhas antigas do dia 24/09 não foram removidas — achei {penultimo_id} onde esperava 3001"
    print("  OK: linhas antigas do dia removidas e substituídas pelas novas, fórmulas arrastadas certo.\n")

    print("\nTODOS OS TESTES PASSARAM.")


if __name__ == "__main__":
    main()
